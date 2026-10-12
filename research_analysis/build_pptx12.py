"""Phase 12 PPTX: 25 slides from FINAL_BTECH_PRESENTATION.md + figures."""
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

slides = [
    ('Title', 'Data-Driven Road Traffic Accident Analysis in India\nMaster title + candidate/guide/institute/2026', None),
    ('Problem', 'Counts without exposure; fragmented vintages; no microdata.', None),
    ('Motivation', 'Evidence-linked analytics + reusable upload platform.', None),
    ('Research Question', 'RQ1-RQ11 (observational framing).', None),
    ('Objectives', '13 objectives (trends, geography, weather, causes, validation, ML, clustering, limits).', None),
    ('Research Framework', 'Discover-Veriy-Preserve-Audit-Harmonize-Analyze-Validate-Model-Synthesize.', None),
    ('Data Sources', 'Benchmark 254 obs; history conditional; NCRB separate; IMD 236 matched; 7 modules; TN.', None),
    ('Data Validation', '0 benchmark mismatches; 12/12 history; NCRB sections exact; TN sums exact.', None),
    ('India-Level Trends', '+3.7% acc / +12.4% fat; 2020 dip; slopes ns.', 'F01'),
    ('Geographical Analysis', '38 entities; heterogeneity; no rankings.', 'F05'),
    ('Zone Analysis', 'Friedman p=3.0e-4 descriptive; zone trends.', 'F04'),
    ('State Analysis', 'Slopes mostly non-detectable (N=7); CIs wide.', 'F06'),
    ('Temporal Analysis', 'Annual fluctuation; YoY table.', 'F02'),
    ('Seasonal/NCRB', 'May 8.84% / Jul 7.46% (NCRB vintage).', 'F13'),
    ('Weather Analysis', 'Rain -0.234/-0.376; Tmax +0.452/+0.498; never causal.', 'F07'),
    ('Severity Analysis', '33.50 to 36.33 per 100 accidents.', 'F03'),
    ('Violations/Causes', 'Over-speeding 70.3% of deaths (reported category).', 'F09'),
    ('Collision/Road-User', 'Hit from Back leads; two-wheelers 46.2% (counts).', 'F11'),
    ('Licence/Safety', 'Valid licence 68.2%; no-helmet 54,122 toll (reported).', 'F12'),
    ('Statistical Findings', 'OLS/CI, Pearson+Spearman, chi2, Friedman, BH-FDR.', None),
    ('ML Methodology', 'Lag-only; train 2019-22 N=141; test 2023-24 N=72; naive baseline.', None),
    ('ML Results', 'Naive RMSE 1038.6/286.3 beats ALL trained (honest null).', 'F15'),
    ('K-Means', 'K=2, silhouette 0.356, descriptive labels.', 'F16'),
    ('Limitations + Contributions', 'Vehicle stock, microdata, N=213; 10 contributions.', None),
    ('Conclusion + Vision', 'Counts are not risk; associations are not causes; analytical platform.', None),
]
for title, body, fig in slides:
    s = prs.slides.add_slide(blank)
    tx = s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(1)).text_frame
    tx.text = title
    bx = s.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(5), Inches(5)).text_frame
    bx.text = body
    if fig:
        try:
            s.shapes.add_picture(f'FINAL_SUBMISSION_PACKAGE/FIGURES/{fig}.png',
                                 Inches(6), Inches(1.2), width=Inches(6.5))
        except Exception as e:
            print(fig, 'missing:', e)
prs.save('FINAL_SUBMISSION_PACKAGE/06_PRESENTATION/FINAL_BTECH_PRESENTATION.pptx')
print('pptx saved:', len(prs.slides), 'slides')
