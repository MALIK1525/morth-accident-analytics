"""Step 1: full B.Tech dissertation DOCX (formatting only; content from Phase 11)."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pandas as pd

PKG = 'FINAL_SUBMISSION_PACKAGE'
doc = Document()

# base: Times New Roman 12, A4, 1in margins
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(1)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.5
for h, s in [('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)]:
    st = doc.styles[h]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(s)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
cap = doc.styles.add_style('CaptionX', 1)
cap.base_style = doc.styles['Normal']
cap.font.size = Pt(10)
cap.font.italic = True


def P(t, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    if align:
        p.alignment = align
    return p


def H(t, l=1):
    return doc.add_heading(t, level=l)


def field(p, instr):
    r = p.add_run()
    f1 = OxmlElement('w:fldChar')
    f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'end')
    r._r.append(f1)
    r._r.append(it)
    r._r.append(f2)


def footer_pgnum(section, restart=None):
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(fp, 'PAGE')
    if restart is not None:
        pg = OxmlElement('w:pgNumType')
        pg.set(qn('w:start'), str(restart))
        section._sectPr.append(pg)


def md_section(path, skip_title=1):
    lines = open(path, encoding='utf-8').read().split('\n')
    for ln in lines[skip_title:]:
        s = ln.strip()
        if not s:
            continue
        if s.startswith('### '):
            H(s[4:], 3)
        elif s.startswith('## '):
            H(s[3:], 2)
        elif s.startswith('# '):
            H(s[2:], 1)
        elif s.startswith('- ') or s.startswith('* '):
            doc.add_paragraph(s[2:], style='List Bullet')
        else:
            P(s)


def csv_table(path, max_rows=30, fontsize=9):
    df = pd.read_csv(path)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = 'Light Grid Accent 1'
    for j, c in enumerate(df.columns):
        cell = t.rows[0].cells[j]
        cell.text = ''
        run = cell.paragraphs[0].add_run(str(c))
        run.bold = True
        run.font.size = Pt(fontsize)
    for _, row in df.head(max_rows).iterrows():
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ''
            run = cells[j].paragraphs[0].add_run('' if pd.isna(v) else str(v))
            run.font.size = Pt(fontsize)
    doc.add_paragraph()


def fig(fid):
    caps = {}
    for line in open(f'{PKG}/05_FIGURES/FINAL_FIGURE_CAPTIONS.md', encoding='utf-8'):
        if line.startswith('F') and '.' in line[:6]:
            caps[line.split('.')[0]] = line.strip()
    p = doc.add_paragraph(style='CaptionX')
    p.add_run(f'Figure {fid}. ' + caps.get(fid, '')).italic = True
    doc.add_picture(f'{PKG}/FIGURES/{fid}.png', width=Inches(5.8))


C = WD_ALIGN_PARAGRAPH.CENTER
P('A Multi-Level Data-Driven Analysis of Road Traffic Accidents in India:', bold=True, align=C)
P('Geographical, Temporal, Seasonal, Environmental and Severity Assessment', bold=True, align=C)
P('B.Tech Research Report — 2026', align=C)
P('[Candidate Name / Roll No. / Department / Institute — to be completed]', align=C)
P('LPU-specific official template was not available in the supplied project materials; conventional academic formatting was applied.', italic=True, align=C)
footer_pgnum(sec)
doc.add_page_break()
H('Certificate'); P('[Guide certificate placeholder.]')
H('Declaration'); P('[Candidate declaration placeholder.]')
H('Acknowledgement'); P('[Candidate to complete: guide, department, data publishers.]')
H('Abstract')
for para in open(f'{PKG}/03_ABSTRACT/FINAL_ABSTRACT.md', encoding='utf-8').read().split('\n\n')[1:]:
    P(para.replace('\n', ' '))
H('Table of Contents')
p = doc.add_paragraph()
field(p, 'TOC \\o "1-3" \\h \\z \\u')
H('List of Figures')
p = doc.add_paragraph()
field(p, 'TOC \\h \\z \\c "Figure"')
H('List of Tables')
p = doc.add_paragraph()
field(p, 'TOC \\h \\z \\c "Table"')
H('Abbreviations')
P('MoRTH — Ministry of Road Transport and Highways; NCRB — National Crime Records Bureau; ADSI — Accidental Deaths and Suicides in India; IMD — India Meteorological Department; CMPG — Climate Monitoring and Prediction Group; ABDB — Administrative Boundary Database; RGI — Registrar General of India; BRSI — Basic Road Statistics of India; TRW — Transport Research Wing; OLS — Ordinary Least Squares; FDR — False Discovery Rate; BH — Benjamini-Hochberg; RMSE — Root Mean Square Error; MAE — Mean Absolute Error; KNN — K-Nearest Neighbours (supervised); NV — Not Verified.')

# main matter: restart arabic numbering
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
footer_pgnum(sec2, restart=1)

H('Chapter 1 — Introduction')
P('Background: India recorded 487,707 accidents and 177,175 deaths in 2024 (+3.7%/+12.4% since 2018); severity rose 33.50 to 36.33 deaths per 100 accidents.')
P('Problem: counts without exposure denominators; fragmented vintages; no microdata. Research gap: reconciliation, geography, and missing-denominator work rarely combined. Research question (observational): how do geographical, temporal, seasonal, environmental, roadway, vehicle and other available factors associate with occurrence and severity?')
P('Objectives 1-13 per Phase 11 (trends; geography; temporal; weather; violations; collisions; road-users; licence/safety; validation; prediction-vs-baseline; clustering; limitations; reusable framework). Scope: 2018-2024 benchmark; conditional history; separate NCRB/IMD/supporting layers.')

H('Chapter 2 — Literature Review')
P('LITERATURE REVIEW SOURCES REQUIRED — no external literature was collected in this package; no citations invented.')

H('Chapter 3 — Research Methodology')
md_section('research_analysis/PHASE10_METHODOLOGY.md')

H('Chapter 4 — Data Sources and Validation')
P('Benchmark: 254 state-years, 38 entities, 0 mismatches. NCRB monthly separate (gap -19,740 vs MoRTH 2024). IMD 236 matched/18 excluded. 7 supporting modules VERIFIED-MIRROR. TN districts reconciled. AR2017 history conditional (12/12 MATCH). RGI projections (not census); BRSI 2021-22 only.')

H('Chapter 5 — India-Level Analysis')
P('Accidents 470,403 to 487,707 (+3.7%) with 2020 dip -18.6%; fatalities 157,593 to 177,175 (+12.4%); OLS slopes ns (N=7).')
for f in ['F01', 'F02', 'F03']:
    fig(f)

H('Chapter 6 — Geographical Analysis')
P('38 entities; six zones; Friedman chi2=25.29, p=3.0e-4, W=0.72 (descriptive differences, no ranking). State slopes mostly non-detectable.')
for f in ['F04', 'F05', 'F06']:
    fig(f)

H('Chapter 7 — Temporal and Seasonal Analysis')
P('Annual fluctuation with 2020 structural break; NCRB 2024 monthly peak May 8.84%, trough July 7.46% (NCRB vintage, never merged to MoRTH).')
fig('F13')

H('Chapter 8 — Environmental Analysis')
P('Pooled rain r=-0.234/-0.376; Tmax +0.452/+0.498; Tmin +0.366/+0.340 (FDR-held, N=236). Scale confounding disclosed; severity null; never causal.')
for f in ['F07', 'F08']:
    fig(f)

H('Chapter 9 — Severity, Cause, Collision and Road-User Analysis')
P('Over-speeding 70.7% accidents / 70.3% fatalities (reported category). Hit from Back most frequent; Hit and Run ratio ~0.50. Two-wheelers 46.2%, pedestrians 20.6% (counts, not risk). Valid licence 68.2% of reported accidents. No-helmet 54,122 / no-seatbelt 14,466 deaths (toll).')
for f in ['F09', 'F10', 'F11', 'F12']:
    fig(f)
fig('F14')

H('Chapter 10 — Statistical Analysis')
P('OLS with CI; Pearson+Spearman; chi-square monthly (1276, large-N); Friedman; BH-FDR; effect sizes. Near-zero rule |slope|<=0.5% of mean (documented).')

H('Chapter 11 — Machine Learning and Clustering')
P('Lag-only features, train 2019-22 (N=141), test 2023-24 (N=72), naive baseline RMSE 1038.6/286.3 beats ALL trained models (honest null; GBR train R2~1.0 = overfitting). KNN supervised. K-means K=2, silhouette 0.356, descriptive labels, ARI stability.')
for f in ['F15', 'F16']:
    fig(f)

H('Chapter 12 — Discussion')
md_section('research_analysis/PHASE10_DISCUSSION_SECTION.md')
H('Chapter 13 — Limitations')
P('Vehicle stock, microdata/coords, driver/vehicle detail, humidity etc., monthly MoRTH, single-state districts, 2023-24 state injuries, conditional history, N=213 panel. Full matrix: PHASE10_LIMITATIONS_MATRIX.csv.')
H('Chapter 14 — Contributions')
md_section('research_analysis/PHASE10_CONTRIBUTIONS.md')
H('Chapter 15 — Application Vision')
P('Mode A/B upload analytics platform; analytical tool, explicitly NOT a proven prediction system.')
H('Chapter 16 — Conclusion')
md_section('research_analysis/PHASE10_CONCLUSION.md')

H('References')
md_section('FINAL_SUBMISSION_PACKAGE/10_REFERENCES/FINAL_REFERENCES.md')

H('Appendices')
P('Appendix A — Data dictionary (benchmark DATA_DICTIONARY + supporting schemas). Appendix B — MASTER_PARAMETER_AVAILABILITY.csv (52 params). Appendix C — Phase 7/8 statistical CSVs. Appendix D — PHASE9_MODEL_RESULTS/PREDICTIONS/BASELINE_COMPARISON. Appendix E — K_SELECTION/RESULTS/PROFILES. Appendix F — graph plans. Appendix G — seeds/scripts/PHASE10_REPRODUCIBILITY_AUDIT.csv.')

H('Final Tables 1-15 (values from FINAL_TABLES.md; full detail in file)')
P(open(f'{PKG}/04_TABLES/FINAL_TABLES.md', encoding='utf-8').read())

doc.save(f'{PKG}/01_FINAL_REPORT/FINAL_BTECH_RESEARCH_REPORT.docx')
print('dissertation docx saved')
